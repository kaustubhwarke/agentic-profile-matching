"""Resume ingestion pipeline.

Reads resume files (``.txt``, ``.md``, ``.pdf``, ``.docx``) from the configured
resume directory, normalises them into LangChain ``Document`` objects with
stable metadata, and indexes them into the vector store.
"""

from __future__ import annotations

from pathlib import Path

from langchain_core.documents import Document

from profile_matching.config import Settings, get_settings
from profile_matching.logging_config import get_logger
from profile_matching.rag.vector_store import ResumeVectorStore, get_vector_store

logger = get_logger(__name__)

_SUPPORTED_SUFFIXES = {".txt", ".md", ".pdf", ".docx"}


def _read_text(path: Path) -> str:
    """Extract plain text from a resume file based on its suffix."""
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md"}:
        return path.read_text(encoding="utf-8", errors="ignore")
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    if suffix == ".docx":
        import docx

        document = docx.Document(str(path))
        return "\n".join(p.text for p in document.paragraphs)
    raise ValueError(f"Unsupported resume format: {path.suffix}")


def _derive_name(text: str, fallback: str) -> str:
    """Best-effort extraction of a candidate name from the first content line."""
    for line in text.splitlines():
        stripped = line.strip()
        if stripped:
            # First non-empty line is treated as the name (common resume layout).
            return stripped[:80]
    return fallback


def load_resume_documents(settings: Settings | None = None) -> list[Document]:
    """Load all resumes from the configured directory into ``Document`` objects."""
    settings = settings or get_settings()
    resume_dir = settings.resume_dir
    if not resume_dir.exists():
        logger.warning("Resume directory %s does not exist.", resume_dir)
        return []

    documents: list[Document] = []
    for path in sorted(resume_dir.iterdir()):
        if path.suffix.lower() not in _SUPPORTED_SUFFIXES:
            continue
        try:
            text = _read_text(path).strip()
        except Exception as exc:  # pragma: no cover - defensive
            logger.error("Failed to read %s: %s", path.name, exc)
            continue
        if not text:
            logger.warning("Skipping empty resume: %s", path.name)
            continue
        documents.append(
            Document(
                page_content=text,
                metadata={
                    "candidate_id": path.stem,
                    "name": _derive_name(text, path.stem),
                    "source": path.name,
                },
            )
        )
    logger.info("Loaded %d resume document(s) from %s.", len(documents), resume_dir)
    return documents


def ingest_resumes(
    settings: Settings | None = None,
    store: ResumeVectorStore | None = None,
) -> int:
    """Load and index all resumes. Returns the number indexed."""
    settings = settings or get_settings()
    store = store or get_vector_store()
    documents = load_resume_documents(settings)
    return store.add_documents(documents)
